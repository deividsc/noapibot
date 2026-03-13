"""
File operation handlers: /read, /save, /pdf, /docx, /pptx, /xlsx and their manual pages.
"""
import re
import asyncio
import tempfile
from datetime import datetime
from pathlib import Path

from telegram import Update
from telegram.ext import ContextTypes

from noapibot.config import OUTPUTS_DIR, SKILLS_DIR
from noapibot.memory import load_memory, save_memory
from noapibot.core import run_opencode, run_with_context, send_response
from noapibot.websocket import broadcast_status
import noapibot.state as state
import noapibot.memory as memory_mod


# ─── Extractors ───────────────────────────────────────

def extract_docx_text(path):
    try:
        from docx import Document
        doc = Document(path)
        content = []
        for para in doc.paragraphs:
            if para.text.strip():
                content.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_data:
                    content.append(" | ".join(row_data))
        return "\n".join(content)
    except Exception as e:
        return f"Error extracting DOCX: {e}"


def extract_pptx_text(path):
    try:
        from pptx import Presentation
        prs = Presentation(path)
        content = []
        for i, slide in enumerate(prs.slides):
            slide_text = [f"--- Diapositiva {i+1} ---"]
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text.strip())
            content.append("\n".join(slide_text))
        return "\n\n".join(content)
    except Exception as e:
        return f"Error extracting PPTX: {e}"


def extract_xlsx_text(path):
    try:
        import pandas as pd
        all_sheets = pd.read_excel(path, sheet_name=None)
        content = []
        for sheet_name, df in all_sheets.items():
            content.append(f"=== HOJA: {sheet_name} ===")
            content.append(df.to_csv(index=False, sep="|"))
            content.append("")
        return "\n".join(content)
    except Exception as e:
        return f"Error extracting XLSX (pandas): {e}"


# ─── File Commands ────────────────────────────────────

async def read_file(update: Update, context: ContextTypes.DEFAULT_TYPE):
    path = ' '.join(context.args) if context.args else None
    if not path:
        await update.message.reply_text("Uso: /read <ruta>\nEj: /read D:\\proyecto\\main.py")
        return
    p = Path(path)
    if not p.exists():
        await update.message.reply_text(f"❌ Archivo no encontrado: {path}")
        return
    try:
        content = p.read_text(encoding='utf-8', errors='replace')
        if len(content) > 4000:
            await update.message.reply_document(document=open(p, 'rb'), caption=f"📂 {p.name} ({len(content)} chars)")
        else:
            await update.message.reply_text(f"📂 `{p.name}`:\n```\n{content}\n```")
    except Exception as e:
        await update.message.reply_text(f"❌ Error: {e}")


async def save_file(update: Update, context: ContextTypes.DEFAULT_TYPE):
    filename = ' '.join(context.args) if context.args else None
    memory = load_memory()
    last_response = None
    for m in reversed(memory):
        if m["role"] == "assistant":
            last_response = m["text"]
            break
    if not last_response:
        await update.message.reply_text("❌ No hay respuesta previa para guardar.")
        return
    if not filename:
        filename = f"output_{datetime.now().strftime('%Y%m%d_%H%M%S')}.md"
    content = last_response
    code_exts = {'.py', '.js', '.ts', '.sh', '.bash', '.java', '.cpp', '.c', '.go', '.rs', '.rb', '.php', '.sql', '.css', '.html'}
    if Path(filename).suffix.lower() in code_exts:
        blocks = re.findall(r'```\w*\n(.*?)```', content, re.DOTALL)
        if blocks:
            content = '\n\n'.join(b.strip() for b in blocks)
    filepath = OUTPUTS_DIR / filename
    filepath.write_text(content, encoding='utf-8')
    await update.message.reply_text(f"💾 Guardado en: `{filepath}`")


# ─── PDF Processing ───────────────────────────────────

async def pdf_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    msg = update.message
    doc = msg.document if msg.document else (msg.reply_to_message.document if msg.reply_to_message and msg.reply_to_message.document else None)
    if not doc or not doc.file_name.lower().endswith('.pdf'):
        await msg.reply_text("📄 Adjunta un PDF o responde a un mensaje con PDF.\nPara ayuda: /pdfs")
        return
    status = await msg.reply_text("📄 Procesando PDF...")
    try:
        file = await doc.get_file()
        tmp = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
        await file.download_to_drive(tmp.name)

        import pdfplumber
        text_parts = []
        with pdfplumber.open(tmp.name) as pdf:
            for i, page in enumerate(pdf.pages):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(f"--- Página {i+1} ---\n{page_text}")
        
        Path(tmp.name).unlink(missing_ok=True)
        
        if not text_parts:
            await status.edit_text("❌ No se pudo extraer texto. El PDF podría ser una imagen.")
            return
        
        full_text = "\n\n".join(text_parts)
        if len(full_text) > 15000:
            full_text = full_text[:15000] + "\n\n... [Truncado a 15000 chars]"
        
        prompt = ' '.join(context.args) if context.args else "Analiza este documento y hazme un resumen completo."
        user_msg = f"[DOCUMENTO PDF: {doc.file_name}]\n\n{full_text}\n\nPETICIÓN: {prompt}"
        
        thread_id = getattr(msg, 'message_thread_id', None)
        session_id = f"topic_{thread_id}" if getattr(msg, 'is_topic_message', False) and thread_id else memory_mod.current_session
        
        result = await run_with_context(state.current_model, user_msg, session_id=session_id, chat_id=update.effective_chat.id, thread_id=thread_id)
        
        memory = load_memory(session_id)
        memory.append({"role": "user", "text": f"[PDF: {doc.file_name}] {prompt}", "ts": datetime.now().isoformat()})
        memory.append({"role": "assistant", "text": result, "ts": datetime.now().isoformat()})
        save_memory(memory, session_id)
        
        await send_response(update, status, result)
        await broadcast_status("idle", "", agent="NoApiBot")
    except Exception as e:
        await status.edit_text(f"❌ Error procesando PDF: {e}")


# ─── Office Processing ────────────────────────────────

async def build_office_cmd(update, context, ext, mime_type, extractor_func):
    msg = update.message
    doc = msg.document if msg.document else (msg.reply_to_message.document if msg.reply_to_message and msg.reply_to_message.document else None)
    if not doc or not doc.file_name.lower().endswith(ext):
        await msg.reply_text(f"📄 Adjunta un archivo {ext.upper()} o responde a un mensaje con uno.")
        return
    status = await msg.reply_text(f"📄 Procesando {ext.upper()}...")
    try:
        file = await doc.get_file()
        tmp = tempfile.NamedTemporaryFile(suffix=ext, delete=False)
        await file.download_to_drive(tmp.name)
        content = await asyncio.to_thread(extractor_func, tmp.name)
        Path(tmp.name).unlink(missing_ok=True)
        if not content or content.startswith("Error"):
            await status.edit_text(f"❌ {content}")
            return
        if len(content) > 15000:
            content = content[:15000] + "\n\n... [Truncado]"
        prompt = ' '.join(context.args) if context.args else "Analiza este documento y hazme un resumen completo."
        user_msg = f"[DOCUMENTO {ext.upper()}: {doc.file_name}]\n\n{content}\n\nPETICIÓN: {prompt}"
        thread_id = getattr(msg, 'message_thread_id', None)
        session_id = f"topic_{thread_id}" if getattr(msg, 'is_topic_message', False) and thread_id else memory_mod.current_session
        result = await run_with_context(state.current_model, user_msg, session_id=session_id, chat_id=update.effective_chat.id, thread_id=thread_id)
        memory = load_memory(session_id)
        memory.append({"role": "user", "text": f"[{ext.upper()}: {doc.file_name}] {prompt}", "ts": datetime.now().isoformat()})
        memory.append({"role": "assistant", "text": result, "ts": datetime.now().isoformat()})
        save_memory(memory, session_id)
        await send_response(update, status, result)
        await broadcast_status("idle", "", agent="NoApiBot")
    except Exception as e:
        await status.edit_text(f"❌ Error: {e}")

async def docx_cmd(u, c): await build_office_cmd(u, c, '.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', extract_docx_text)
async def pptx_cmd(u, c): await build_office_cmd(u, c, '.pptx', 'application/vnd.openxmlformats-officedocument.presentationml.presentation', extract_pptx_text)
async def xlsx_cmd(u, c): await build_office_cmd(u, c, '.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', extract_xlsx_text)

# ─── Manual Pages ─────────────────────────────────────

async def pdfs_cmd(u, c): await u.message.reply_text("📄 /pdf – Envía un PDF adjunto o responde a un mensaje con PDF y usa `/pdf <pregunta>`")
async def docxs_cmd(u, c): await u.message.reply_text("📄 /docx – Envía un Word adjunto o responde a un mensaje con .docx y usa `/docx <pregunta>`")
async def pptxs_cmd(u, c): await u.message.reply_text("📄 /pptx – Envía un PowerPoint adjunto o responde a un mensaje con .pptx y usa `/pptx <pregunta>`")
async def xlsxs_cmd(u, c): await u.message.reply_text("📄 /xlsx – Envía un Excel adjunto o responde a un mensaje con .xlsx y usa `/xlsx <pregunta>`")
